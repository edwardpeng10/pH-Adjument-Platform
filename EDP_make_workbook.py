import os
import re
import sys
import json
import ntpath
import shutil
from datetime import datetime
import pandas as pd
# EDP_VERSION = "EDP-v17-2026-10"   (EDP_tecan_gp_master.py Check: all six EDP files must carry the same version)

if sys.stdout is None: sys.stdout = open(os.devnull, "w")      # started with pythonw: there is no console
if sys.stderr is None: sys.stderr = open(os.devnull, "w")

# Excel output at the end of every run.
#   EDP_make_workbook.py
# 1. v16: THE RUN'S OWN WORKBOOK, for taking the data off the robot PC (Edward, 22 Oct: titration_data.xlsx could not be
#    exported from the computers). One small file per run, values only (no formulas, links or macros), named after the
#    batch and the run's start: EDP_results_<batch>_<yyyymmdd_hhmmss>.xlsx, written to EXPORT_DIRS - v17 (Edward, 22 Oct:
#    the Public Documents copy could not be exported either): Edward's folder on the network H: drive first (by its
#    letter; if this process cannot see the letter, by the network name the user's registry gives for it), then Public
#    Documents\EDP_results if H: cannot be reached (not mapped for the account EVOware runs Python as, or offline), then
#    ML_data_archive\EDP_results; the folder used is logged and given to the final message (workbook_location.csv:
#    1 = H:, 2 = Public Documents, 3 = ML_data_archive, 0 = not written) - and copied into ML_data_archive\runs\<Run_ID>\. Sheets: Run (the
#    numbers of the two RATIOS STORED messages, the probe drift in pH and %, robot and waiting time, the chemistry
#    fingerprint), Wells (every well and lane read: plate, well, step, contents per mL eluate, pH typed), Stored ratios,
#    Probe drift (each column-2 well read at plate checks 2 and 3: change in pH and in % of the reading), then this run's rows of every archive table (Readings, Targets, Titrations, RatioMap, Pools,
#    Chemistry, Warnings), Timing and the python log. Each run's rows have the archive's own columns, so the run files
#    put together rebuild the whole database.
# 2. titration_data.xlsx in ML_data_archive, as before: every archive table as one sheet, rewritten each run (the CSVs
#    stay the source of truth). v15: if it is open in Excel, titration_data_<date_time>.xlsx is written instead.
ARCHIVE_DIR = os.environ.get("TECAN_ARCHIVE_DIR", r"C:\ProgramData\Tecan\EVOware\database\ML_data_archive")
WORKING_DIR = os.environ.get("TECAN_DATA_DIR", r"C:\ProgramData\Tecan\EVOware\database\tecan_data")
EXPORT_DIRS = [os.environ.get("EDP_EXPORT_DIR", r"H:\Instrument_Outputs\All_Scientists_Storage\kplc540 - Edward Peng"),
               r"C:\Users\Public\Documents\EDP_results"]       # v17: in this order, then ML_data_archive\EDP_results
SHEETS = [
    ("Runs", "sessions.csv", "One row per run: batch, product, titrants, entries typed, duration, robot and waiting time, probe drift."),
    ("Titrations", "experiment_summary.csv", "One row per titration: conditions, status, metadata, Cycle-2 plan."),
    ("Targets", "target_results.csv", "Each target: ratio used, recommended volume, validated pH, error, refit ratio."),
    ("Predictions", "predictions.csv", "What the chemistry predicted from Cycle 1 alone, against the measured answer."),
    ("Chemistry", "chemistry_fits.csv", "Fitted chemistry per run: the parameters for a predictive model."),
    ("Wells", "wells.csv", "v17: every well and lane read in every run - the ML table: acid and Tris per mL eluate (mmol), the Tris stock, the well's total volume per mL eluate, target, pH typed, left out by the fit."),
    ("Curves", "measurements.csv", "Every curve point with its cycle, volume, ratio and pH."),
    ("Pools", "pools.csv", "Each pool made in the 6-well plate: lane, target, route (acid, or Tris since v13), volumes, pH after each round, trims, final error, Tris ratios."),
    ("RatioMap", "ratio_map.csv", "v14+: per run, mL titrant per mL to go from the eluate (acid, base paths) or from any acid-path pH (neutralise) to each target pH; from pH A to B on one path: (ratio B - ratio A) / (1 + ratio A). Model_Per_mL_eluate: the chemistry model's value, an independent check."),
    ("Warnings", "experiment_warnings.csv", "Anything flagged during a run."),
]
RUN_TABLES = [("Readings", "measurements.csv"), ("Targets", "target_results.csv"), ("Titrations", "experiment_summary.csv"),
              ("RatioMap", "ratio_map.csv"), ("Pools", "pools.csv"), ("Chemistry", "chemistry_fits.csv"),
              ("Warnings", "experiment_warnings.csv")]


def log(msg):
    try:
        with open(os.path.join(WORKING_DIR, "python_log.txt"), "a") as f:
            f.write(msg + "\n")
    except Exception:
        pass


def network_name(path):
    """v17: the same folder by its network name (\\\\server\\share\\...) when its drive letter cannot be seen by this
    process - Windows hides a user's mapped drives from programs started as administrator, which EVOware may be. The
    user's mapped drives are kept in the registry (HKEY_CURRENT_USER\\Network\\<letter>: RemotePath). None if the
    letter can be seen, or no mapping is found."""
    drive, rest = ntpath.splitdrive(path)
    if len(drive) != 2 or drive[1] != ":" or os.path.isdir(drive + "\\"):
        return None
    try:
        import winreg
        with winreg.OpenKey(winreg.HKEY_CURRENT_USER, "Network\\" + drive[0].upper()) as key:
            remote = str(winreg.QueryValueEx(key, "RemotePath")[0]).rstrip("\\")
        return remote + rest if remote.startswith("\\\\") else None
    except Exception:
        return None


def text(name, folder=WORKING_DIR):
    try:
        with open(os.path.join(folder, name), encoding="latin-1") as f:
            return f.read()
    except Exception:
        return ""


def numbers(name):
    out = []
    for tok in re.split(r"[,;\s]+", text(name).replace('"', " ")):
        try:
            out.append(float(tok))
        except ValueError:
            pass
    return out


def table(name, folder=ARCHIVE_DIR):
    try:
        return pd.read_csv(os.path.join(folder, name))
    except Exception:
        return pd.DataFrame()


def writable(path):
    """False if the file exists and Windows refuses to overwrite it (open in Excel)."""
    try:
        if os.path.exists(path):
            with open(path, "a+b"):
                pass
        return True
    except PermissionError:
        return False


def style(writer):
    from openpyxl.styles import Font
    for ws in writer.book.worksheets:
        ws.freeze_panes = "A2"
        for cell in ws[1]:
            cell.font = Font(name="Arial", bold=True)
        for column in ws.columns:
            width = max((len(str(c.value)) for c in column if c.value is not None), default=10)
            ws.column_dimensions[column[0].column_letter].width = min(max(width + 2, 10), 60)


# ----------------------------------------------------------------------------- 1. the run's own workbook (v16)
def run_sheet(run_id, batch, st, summ, opts, drift, sess, chem):
    """Item / value / note rows: what the run stored and how it went."""
    rows = []
    add = lambda item, value, note="": rows.append({"Item": item, "Value": value, "Note": note})
    get = lambda i: summ[i] if len(summ) > i else None
    ratio = lambda x: None if x is None or x < 0 else round(x, 4)
    started = ""
    if run_id.startswith("RUN_"):
        try:
            started = datetime.strptime(run_id[4:], "%Y%m%d_%H%M%S").strftime("%Y-%m-%d %H:%M:%S")
        except ValueError:
            pass
    add("Run ID", run_id); add("Started", started)
    add("Batch ID", batch[0]); add("Product name", batch[1]); add("Product type", batch[2])
    add("Analysis version", st.get("version", ""))
    add("Tips 5-8 used", "yes" if not (len(opts) > 2 and opts[2] == 0) else "no", "Use_8_tips in the script (v16)")
    add("VI neutralisation titration", "yes" if opts[:1] == [1.0] else "no")
    add("Base titration of the eluate", "yes" if len(opts) > 1 and opts[1] == 1 else "no")
    dil = (st.get("base") or {}).get("dil")
    if len(opts) > 1 and opts[1] == 1:
        add("Dilute Tris for the base wells, M", dil if dil else "not used",
            "Product 2W RIGHT well; each base well that needs at most 980 uL of it (Wells sheet: Tris_Stock_M) - v17")
    a = st.get("acid", {})
    add("Eluate pH", a.get("ph0"), "A1, product only")
    if drift:
        add("Probe drift, pH", drift["mean_ph"], "column 2 read again at plate check 3 minus its plate-check-2 reading, mean of the wells")
        add("Probe drift, %", drift["mean_pct"], "the same change as % of the reading")
        add("Probe drift taken off the acid validation", drift["applied_ph"], f"only when the mean change is {drift['limit']} pH or more")
    else:
        add("Probe drift", "not measured", "column 2 not read again at plate check 3")
    tp = numbers("titration_params.csv")
    targets = tp[2:6] if len(tp) >= 6 else [3.6, 3.4, 3.2, 3.5]
    names = ("target 1", "target 2", "target 3", "process")
    for k, (t, nm) in enumerate(zip(targets, names)):
        add(f"Acid ratio to pH {t:g} ({nm})", ratio(get(1 + k)), "mL 1 M acetic acid per mL eluate, after the validation refit")
    add("Acid: worst validation error", get(0), "pH")
    bt = (numbers("tris_params.csv") or [7.2, 7.4, 7.6, 7.4])[:4]
    for k, t in enumerate(bt[:3]):
        v = get(11 + k)
        add(f"Base ratio to pH {t:g}", ratio(v) if v not in (None, 0) else ("at or below the eluate pH" if v == 0 else None),
            "mL 1 M Tris per mL eluate, Tris straight into the eluate")
    add("Base: worst validation error", get(10), "pH (-1 = not run)")
    proc = bt[3] if len(bt) > 3 else 7.4
    for k, (t, nm) in enumerate(zip(targets, names)):
        add(f"Neutralisation from pH {t:g} ({nm}) to pH {proc:g}", ratio(get(15 + k)), "mL 1 M Tris per mL of acidified product")
    for k, t in ((0, bt[0]), (2, bt[2] if len(bt) > 2 else 7.6)):
        add(f"Neutralisation from the process pool to pH {t:g}", ratio(get(6 + k)), "mL 1 M Tris per mL of acidified product")
    add("VI neutralisation: worst validation error", get(5), "pH (-1 = not run)")
    if not sess.empty:
        r = sess.iloc[-1]
        add("Robot minutes", r.get("Robot_Min"), "the robot working")
        add("Waiting minutes", r.get("Stop_Min"), "the robot waiting for someone")
        add("Entries typed", r.get("Entries_Typed"))
    if not chem.empty:
        r = chem.iloc[-1]
        for col, label in (("Acetate_M", "Chemistry model: acetate, M"), ("Strong_Ions_M", "Chemistry model: strong-ion difference, M"),
                           ("Carboxyl_M", "Chemistry model: carboxyls, M"), ("Histidine_M", "Chemistry model: histidine-like groups, M"),
                           ("Tris_In_Eluate_M", "Chemistry model: Tris in the eluate, M"), ("Model_Temp_C", "Chemistry model: temperature, C"),
                           ("Tris_Strength", "Chemistry model: Tris strength relative to the acid"), ("Fit_RMS_pH", "Chemistry model: RMS, pH")):
            if col in r.index:
                v = r.get(col)
                if col == "Tris_Strength" and (v is None or pd.isna(v)):
                    v = 1.0                                  # not fitted: the model used 1:1 stocks
                add(label, v, ("fitted (the VI wells asked for it)" if r.get("Tris_Strength_Fitted") == 1 else "1:1 stocks assumed")
                    if col == "Tris_Strength" else "")
    return pd.DataFrame(rows)


def well_rows(st):
    """Every well and lane read this run: plate, well, step, what it held per mL eluate (acid and Tris in mL of the 1 M
    stocks, i.e. mmol; v17: the Tris stock it was made with and the well's total volume, so the concentrations in the
    well follow), the pH typed, and whether the analysis left that reading out of its fit."""
    rows = []
    a = st.get("acid") or {}
    b = st.get("base") or {}
    tr = st.get("tris") or {}
    left = {(round(x, 5), round(y, 2)) for x, y in (a.get("dropped") or [])} | \
           {("b", round(x, 5), round(y, 2)) for x, y in (b.get("dropped") or [])} | \
           {("v", round(x, 5), round(y, 2)) for x, y in (tr.get("dropped") or [])}
    entered = {}                                         # v17: when each entry was typed (the last accepted one per file)
    try:
        import csv
        with open(os.path.join(WORKING_DIR, "entries_log.csv"), newline="") as f:
            for r in csv.DictReader(f):
                if r.get("accepted") == "1":
                    entered[r.get("file")] = r.get("timestamp")
    except Exception:
        pass
    def add(plate, wells, step, comp, phs, targets=None, stock=None, base=False, fname=None, vi=None):
        for k, (w, (ac, tr)) in enumerate(zip(wells, comp)):
            ph = phs[k] if k < len(phs) else None
            m = stock[k] if stock and k < len(stock) else 1.0
            key = (("v", round(vi[k], 5), round(ph, 2)) if vi and ph else ("b", round(tr, 5), round(ph, 2)) if base and ph
                   else (round(ac, 5), round(ph, 2)) if ph else None)
            rows.append({"Plate": plate, "Well": w, "Step": step, "Acid_mL_per_mL_eluate": round(ac, 5),
                         "Tris_mL_per_mL_eluate": round(tr, 5), "Tris_Stock_M": m if tr else None,
                         "Total_mL_per_mL_eluate": round(1 + ac + (tr / m if tr else 0.0), 5),
                         "Target_pH": targets[k] if targets and k < len(targets) else None,
                         "pH": ph if ph else None, "Left_out_by_fit": int(key in left) if key else None,
                         "Entered": entered.get(fname)})
    col = lambda c: [f"{r}{c}" for r in "ABCD"]
    if a.get("c1"):
        c1 = numbers("acid_c1_measured.csv")
        add("Acid titration", ["A1"], "product only (the eluate pH)", [(0.0, 0.0)], c1[:1], fname="acid_c1_measured.csv")
        add("Acid titration", col(1)[1:], "acid Cycle 1, low acid", [(r, 0.0) for r in a["c1"][:3]], c1[1:4], fname="acid_c1_measured.csv")
        add("Acid titration", col(5), "acid Cycle 1, locate", [(r, 0.0) for r in a["c1"][3:7]], c1[4:8], fname="acid_c1_measured.csv")
    if a.get("c2"):
        add("Acid titration", col(2) + col(3), "acid Cycle 2", [(r, 0.0) for r in a["c2"]], numbers("cycle2_measured.csv"),
            fname="cycle2_measured.csv")
    tp = numbers("titration_params.csv")
    if a.get("ratios"):
        add("Acid titration", col(4), "acid validation", [(r, 0.0) for r in a["ratios"]], numbers("validation_results.csv"),
            tp[2:6] if len(tp) >= 6 else None, fname="validation_results.csv")
    d = a.get("drift_check")
    if d and a.get("c2"):
        add("Acid titration", d["wells"], "drift check (column 2 read again at plate check 3)",
            [(a["c2"]["ABCD".index(w[0])], 0.0) for w in d["wells"]], d["pc3"], fname="validation_results.csv")
    bt = (numbers("tris_params.csv") or [7.2, 7.4, 7.6, 7.4])[:4]
    for key, c, step, fname in (("c1", 4, "base Cycle 1", "base_c1_measured.csv"), ("c2", 5, "base Cycle 2", "base_c2_measured.csv"),
                                ("val_wells", 6, "base validation", "base_val_measured.csv")):
        if b.get(key):
            stock = b.get("val_m" if key == "val_wells" else f"{key}_m")
            add("Base titration", col(c), step, [(0.0, t) for t in b[key]], numbers(fname), bt if key == "val_wells" else None,
                stock, base=True, fname=fname)
    c1 = numbers("tris_c1_measured.csv")
    for k, grp in enumerate(tr.get("c1") or []):
        add("Base titration", col(1 + k), "VI neutralisation Cycle 1", [(grp["rp"], t * (1 + grp["rp"])) for t in grp["t"]],
            c1[4 * k:4 * k + 4], fname="tris_c1_measured.csv", vi=grp["t"])
    if tr.get("val_t"):
        add("Base titration", col(3), "VI neutralisation validation", [(tr["val_rp"], t * (1 + tr["val_rp"])) for t in tr["val_t"]],
            numbers("tris_val_measured.csv"), bt, fname="tris_val_measured.csv", vi=tr["val_t"])
    for pl in st.get("pools") or []:
        e = pl.get("E") or 0.0
        hist = pl.get("hist") or []                          # v17: the pool's ratio when each reading was taken
        rounds = [r for r in (1, 2, 3) if pl["lane"] in (st.get("measure") or {}).get(str(r), [])]   # pool checks it was read at
        for k, ph in enumerate(pl.get("readings") or []):
            if k < len(hist):
                r = hist[k][0]
                ac, trs = (0.0, r) if pl.get("route") == "base" else (r, 0.0)
            else:
                ac = (pl.get("A") or 0.0) / e if e else None
                trs = (pl.get("B") or 0.0) / e if e else None
            tot = (pl["vol"][k] if k < len(pl.get("vol") or []) else 1 + ac + trs if ac is not None else None)
            sol = tot - 1 - ac if tot is not None and ac is not None else 0.0     # mL of Tris solutions per mL eluate
            rows.append({"Plate": "Bulk Pools", "Well": f"lane {pl['lane']}",
                         "Step": f"pool check {rounds[k] if k < len(rounds) else k + 1} ({pl.get('route', '')})",
                         "Acid_mL_per_mL_eluate": round(ac, 5) if ac is not None else None,
                         "Tris_mL_per_mL_eluate": round(trs, 5) if trs is not None else None,
                         # v17: 1 M; less when trims came from the 0.1 M Tris (mmol per mL of the Tris solutions added)
                         "Tris_Stock_M": (round(min(trs / sol, 1.0), 4) if sol > 1e-9 else 1.0) if trs else None,
                         "Total_mL_per_mL_eluate": round(tot, 5) if tot is not None else None,
                         "Target_pH": pl.get("target"), "pH": ph or None, "Left_out_by_fit": None,
                         "Entered": entered.get(f"pools_measured_{rounds[k]}.csv") if k < len(rounds) else None})
    return pd.DataFrame(rows)


def drift_rows(st):
    """The drift check (v15, plate check 3: column 2 read again) per well, in pH and in % of the reading."""
    d = (st.get("acid") or {}).get("drift_check")
    if not d:
        return None, pd.DataFrame()
    rows = [{"Well": w, "Plate check 2 pH": p2, "Plate check 3 pH": p3, "Change pH": round(p3 - p2, 3),
             "Change %": round(100.0 * (p3 - p2) / p2, 2)} for w, p2, p3 in zip(d["wells"], d["pc2"], d["pc3"])]
    rows.append({"Well": "mean", "Plate check 2 pH": None, "Plate check 3 pH": None, "Change pH": d["mean_ph"], "Change %": d["mean_pct"]})
    rows.append({"Well": "taken off the acid validation readings", "Change pH": d["applied_ph"]})
    return d, pd.DataFrame(rows)


def ratio_rows(run_id, summ, drift=None):
    """The stored ratios per titration and target: planned, validated, refit."""
    tr, ex = table("target_results.csv"), table("experiment_summary.csv")
    if tr.empty or "Run_ID" not in tr:
        return pd.DataFrame()
    tr = tr[tr.Run_ID == run_id].copy()
    kind = {}
    if not ex.empty and "Run_ID" in ex:
        for _, r in ex[ex.Run_ID == run_id].iterrows():
            pool = str(r.get("Pool", ""))
            kind[r.Experiment_ID] = {"EL": "Base titration of the eluate (mL Tris per mL eluate)",
                                     "TR": "VI neutralisation at the process pool (mL Tris per mL acidified product)"}.get(
                pool, "Acid titration (mL acetic acid per mL eluate)")
    tr.insert(0, "Titration", [kind.get(e, "") for e in tr.Experiment_ID])
    off = (drift or {}).get("applied_ph") or 0.0         # the acid validation readings as the refit used them
    acid = tr.Titration.str.startswith("Acid")
    tr["Validated_pH_drift_off"] = tr.Validated_pH.where(~acid, tr.Validated_pH - off).round(3)
    keep = ["Titration", "Target_pH", "Is_Process_Target", "Ratio", "Validated_pH", "Validated_pH_drift_off", "Absolute_Error",
            "Corrected_Ratio", "Correction_Pct", "Method"]
    return tr[[c for c in keep if c in tr.columns]].rename(columns={"Ratio": "Planned_Ratio", "Corrected_Ratio": "Stored_Ratio"})


def run_workbook():
    run_id = text("run_id.csv").strip() or "RUN_unknown"
    try:
        st = json.loads(text("run_state.json") or "{}")
    except ValueError:
        st = {}
    raw = [x.strip().strip('"') for x in text("batch_info.csv").strip().split(",")] + ["", "", ""]
    batch = [x if x and x.upper() != "N/A" else "" for x in raw[:3]]
    tag = re.sub(r"[^A-Za-z0-9_-]+", "-", batch[0]).strip("-")[:40] or "nobatch"
    name = f"EDP_results_{tag}_{run_id[4:] if run_id.startswith('RUN_') else datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    where = os.path.join(WORKING_DIR, "workbook_location.csv")     # v17: for the final message (0 = not written)
    try:
        with open(where, "w") as f:
            f.write("0\n")
    except Exception:
        pass
    summ, opts = numbers("run_summary.csv"), numbers("options.csv")
    drift, drift_df = drift_rows(st)
    sess = table("sessions.csv")
    sess = sess[sess.Run_ID == run_id] if "Run_ID" in sess else pd.DataFrame()
    chem = table("chemistry_fits.csv")
    chem = chem[(chem.Run_ID == run_id) & (chem.Phase == "Model")] if {"Run_ID", "Phase"} <= set(chem.columns) else pd.DataFrame()
    sheets = [("Run", run_sheet(run_id, batch, st, summ, opts, drift, sess, chem)), ("Wells", well_rows(st)),
              ("Stored ratios", ratio_rows(run_id, summ, drift)), ("Probe drift", drift_df)]
    for sheet, fname in RUN_TABLES:
        df = table(fname)
        sheets.append((sheet, df[df.Run_ID == run_id] if "Run_ID" in df else pd.DataFrame()))
    timing = table("timing_log.csv", WORKING_DIR)
    sheets.append(("Timing", timing))
    sheets.append(("Log", pd.DataFrame({"python_log.txt": text("python_log.txt").splitlines()})))
    out = None
    unc = network_name(EXPORT_DIRS[0])
    folders = [(1, EXPORT_DIRS[0])] + ([(1, unc)] if unc else []) + [(2, EXPORT_DIRS[1]), (3, os.path.join(ARCHIVE_DIR, "EDP_results"))]
    for i, (k, folder) in enumerate(folders):
        try:                                                 # v17: H: drive (by letter, then by network name), then
            os.makedirs(folder, exist_ok=True)               # Public Documents, then the archive
            path = os.path.join(folder, name)
            if not writable(path):
                path = path.replace(".xlsx", f"_{datetime.now():%H%M%S}.xlsx")
            with pd.ExcelWriter(path, engine="openpyxl") as writer:
                for sheet, df in sheets:
                    (df if not df.empty else pd.DataFrame({"(nothing this run)": []})).to_excel(writer, sheet_name=sheet, index=False)
                style(writer)
            out = path
            with open(where, "w") as f:
                f.write(f"{k}\n")
            break
        except Exception as e:
            log(f"Results workbook: could not write to {folder} ({e!r})" + (" - trying the next folder" if i < len(folders) - 1 else ""))
    if out is None:
        return
    try:                                                     # the archive keeps a copy with the run's working files
        dest = os.path.join(ARCHIVE_DIR, "runs", run_id)
        os.makedirs(dest, exist_ok=True)
        shutil.copyfile(out, os.path.join(dest, os.path.basename(out)))
    except Exception:
        pass
    log(f"Results workbook for this run: {out}")


# ----------------------------------------------------------------------------- v17: every well of every run, one table
def archive_wells():
    """The run's wells (as on its Wells sheet) appended to ML_data_archive\\wells.csv with the run's identity - one table
    of every well and lane ever read, for the ML (Edward, 22 Oct: the point of this stage is accurate data for machine
    learning). A run written again replaces its own rows."""
    run_id = text("run_id.csv").strip() or "RUN_unknown"
    try:
        st = json.loads(text("run_state.json") or "{}")
    except ValueError:
        st = {}
    df = well_rows(st)
    if df.empty:
        return
    raw = [x.strip().strip('"') for x in text("batch_info.csv").strip().split(",")] + ["", "", ""]
    batch = [x if x and x.upper() != "N/A" else "" for x in raw[:3]]
    df.insert(0, "Run_ID", run_id)
    df.insert(1, "Batch_ID", batch[0]); df.insert(2, "Product_Name", batch[1]); df.insert(3, "Product_Type", batch[2])
    df["Algorithm_Version"] = st.get("version", "")
    path = os.path.join(ARCHIVE_DIR, "wells.csv")
    old = table("wells.csv")
    if not old.empty and "Run_ID" in old:
        old = old[old.Run_ID != run_id]
        df = pd.concat([old, df], ignore_index=True)
    os.makedirs(ARCHIVE_DIR, exist_ok=True)
    df.to_csv(path, index=False)


# ----------------------------------------------------------------------------- 2. the whole archive (titration_data.xlsx)
def archive_workbook():
    os.makedirs(ARCHIVE_DIR, exist_ok=True)
    out = os.path.join(ARCHIVE_DIR, "titration_data.xlsx")
    if not writable(out):
        out = os.path.join(ARCHIVE_DIR, f"titration_data_{datetime.now():%Y%m%d_%H%M%S}.xlsx")
        log(f"titration_data.xlsx is open (in Excel?) - this run's copy is {os.path.basename(out)}")
    notes, frames = [], []
    for sheet, name, description in SHEETS:
        df = table(name)
        frames.append((sheet, df))
        notes.append({"Sheet": sheet, "Source file": name, "Rows": len(df), "What it holds": description})
    with pd.ExcelWriter(out, engine="openpyxl") as writer:
        pd.DataFrame(notes).to_excel(writer, sheet_name="Contents", index=False)
        for sheet, df in frames:
            (df if not df.empty else pd.DataFrame({"(no rows yet)": []})).to_excel(writer, sheet_name=sheet, index=False)
        for ws in writer.book.worksheets:
            if ws.max_row > 1 and ws.max_column > 0:
                ws.auto_filter.ref = ws.dimensions
        style(writer)


def main():
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        return log("openpyxl is not installed, so no workbook was built. Install it with: pip install openpyxl")
    for job in (run_workbook, archive_wells, archive_workbook):   # each on its own: one failing never stops the others
        try:
            job()
        except Exception as e:
            log(f"Workbook not built ({job.__name__}): {e!r}")


if __name__ == "__main__":
    try:
        main()
    except Exception as e:                                 # never let it disturb the run
        log(f"Workbook not built: {e!r}")
